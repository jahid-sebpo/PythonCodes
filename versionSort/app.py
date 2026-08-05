import os
import openpyxl

def merge_matrix_to_crm(filename):
    """
    Reads an Excel file from the 'file' directory, matches rows between 'matrix' 
    and 'crm' sheets using Column 2 (B), and copies Columns 3+ from matrix to crm.
    
    Args:
        filename (str): Name of the Excel file (e.g., 'data.xlsx')
    """
    # Construct the file path targeting the folder named 'file'
    folder_name = "file"
    filepath = os.path.join(folder_name, filename)
    
    print(f"Attempting to open: {filepath}")
    
    try:
        # Load the workbook
        workbook = openpyxl.load_workbook(filepath)
        
        # Verify that both expected sheets exist
        required_sheets = ["crm", "matrix"]
        for sheet in required_sheets:
            if sheet not in workbook.sheetnames:
                print(f"Error: Required sheet '{sheet}' was not found in the workbook.")
                print(f"Available sheets: {workbook.sheetnames}")
                return
                
        crm_sheet = workbook["crm"]
        matrix_sheet = workbook["matrix"]
        
        print("Analyzing 'matrix' sheet...")
        # Dictionary to store matrix data. Key: Column 2 value, Value: List of values from Col 3 onwards
        matrix_lookup = {}
        
        # Iterate through 'matrix' starting from row 2 (assuming row 1 is the header)
        for r_idx in range(2, matrix_sheet.max_row + 1):
            key = matrix_sheet.cell(row=r_idx, column=2).value
            
            # Only store if key is not None
            if key is not None:
                row_values = []
                # Grab values from Column 3 to the last column containing data
                for c_idx in range(3, matrix_sheet.max_column + 1):
                    row_values.append(matrix_sheet.cell(row=r_idx, column=c_idx).value)
                matrix_lookup[key] = row_values

        print(f"Found {len(matrix_lookup)} matching keys in 'matrix'. Merging to 'crm'...")
        
        update_count = 0
        # Iterate through 'crm' sheet starting from row 2
        for r_idx in range(2, crm_sheet.max_row + 1):
            crm_key = crm_sheet.cell(row=r_idx, column=2).value
            
            # If the key exists in our matrix lookup, perform the copy
            if crm_key is not None and crm_key in matrix_lookup:
                matrix_values = matrix_lookup[crm_key]
                
                # Write each value starting from Column 3
                for offset, val in enumerate(matrix_values):
                    crm_sheet.cell(row=r_idx, column=3 + offset, value=val)
                update_count += 1

        # Save changes back to the original file
        workbook.save(filepath)
        print(f"Success! {update_count} rows in 'crm' were successfully updated.")
        print(f"Changes saved directly to '{filepath}'")

    except FileNotFoundError:
        print(f"Error: The file '{filename}' was not found inside the directory '{folder_name}/'.")
        print(f"Please ensure '{folder_name}' folder exists and contains '{filename}'.")
    except Exception as e:
        print(f"An unexpected error occurred: {e}")


def filter_crm_by_matrix_versions(
    filename="input.xlsx", 
    crm_header_row=11, 
    crm_version_col=2, 
    matrix_version_col=1, 
    matrix_start_row=1
):
    """
    Reads an Excel file from the 'file' directory.
    - Collects allowed version names from Column A of the 'matrix' sheet.
    - Inspects Column B of the 'crm' sheet starting below header row 11.
    - Keeps CRM rows whose version matches a matrix version; deletes all others.
    """
    folder_name = "file"
    filepath = os.path.join(folder_name, filename)
    
    print(f"\n--- Filtering CRM Tab by Matrix Versions ---")
    print(f"Target File: {filepath}")
    
    if not os.path.exists(filepath):
        print(f"Error: File '{filepath}' not found.")
        return

    def clean_val(val):
        """Converts value to a clean string for accurate matching."""
        if val is None:
            return ""
        s = str(val).strip().replace("\xa0", "")
        if s.endswith(".0"):
            s = s[:-2]
        return s.lower()

    try:
        workbook = openpyxl.load_workbook(filepath)
        
        # Find sheet names regardless of case/spaces
        sheet_map = {name.strip().lower(): name for name in workbook.sheetnames}
        if "crm" not in sheet_map or "matrix" not in sheet_map:
            print(f"Error: Could not find both 'crm' and 'matrix' tabs in workbook.")
            print(f"Found sheets: {workbook.sheetnames}")
            return
            
        crm_sheet = workbook[sheet_map["crm"]]
        matrix_sheet = workbook[sheet_map["matrix"]]

        # 1. Collect all version names from Matrix Column A (Column 1)
        valid_versions = set()
        raw_matrix_versions = []
        for r_idx in range(matrix_start_row, matrix_sheet.max_row + 1):
            val = matrix_sheet.cell(row=r_idx, column=matrix_version_col).value
            cleaned = clean_val(val)
            if cleaned:
                valid_versions.add(cleaned)
                raw_matrix_versions.append(str(val).strip())

        print(f"Found {len(valid_versions)} valid versions in Matrix tab.")
        print(f"Matrix versions sample: {raw_matrix_versions[:5]}")

        # 2. Inspect CRM tab (Data starts on row 12)
        crm_start_row = crm_header_row + 1
        crm_max_row = crm_sheet.max_row

        if crm_max_row < crm_start_row:
            print(f"No data rows found below row {crm_header_row} in CRM tab.")
            return

        # 3. Iterate backwards from bottom row to start row to safely delete non-matching rows
        kept_count = 0
        deleted_count = 0

        for r_idx in range(crm_max_row, crm_start_row - 1, -1):
            crm_version_val = crm_sheet.cell(row=r_idx, column=crm_version_col).value
            cleaned_crm_v = clean_val(crm_version_val)

            # Check if CRM version exists in Matrix valid versions set
            if cleaned_crm_v in valid_versions:
                kept_count += 1
            else:
                crm_sheet.delete_rows(r_idx)
                deleted_count += 1

        # 4. Save workbook directly
        workbook.save(filepath)
        print(f"\nCompleted successfully!")
        print(f"Rows Kept in CRM tab: {kept_count}")
        print(f"Rows Deleted from CRM tab: {deleted_count}")
        print(f"File updated at: '{filepath}'")

    except Exception as e:
        print(f"An error occurred during filtering: {e}")


# --- Main execution block ---
if __name__ == "__main__":
    target_excel_file = "input.xlsx" 
    
    # Run filtering: CRM header at row 11, CRM version in Col B (2), Matrix version in Col A (1)
    filter_crm_by_matrix_versions(
        filename=target_excel_file, 
        crm_header_row=11, 
        crm_version_col=2, 
        matrix_version_col=1, 
        matrix_start_row=1
    )